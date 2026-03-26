"""
Shared tools for the marketing email generation pipeline.

Tools:
    perplexity_search    — Web research via Perplexity Sonar API
    generate_docx_report — Word document generation via python-docx
"""

from __future__ import annotations

import json
import os
from datetime import datetime
from pathlib import Path

import requests
from strands import tool

PPLX_API_KEY: str = os.environ.get("PPLX_API_KEY", "")


# ─────────────────────────────────────────────────────────────────────────────
# Perplexity Search
# ─────────────────────────────────────────────────────────────────────────────

@tool
def perplexity_search(query: str, model: str = "sonar") -> str:
    """Search the web using Perplexity AI for accurate, up-to-date information.

    Use this for ALL research tasks — company data, employee profiles,
    industry trends, geopolitical analysis, product capabilities, etc.

    Args:
        query: The research question or search query. Be specific and detailed.
        model: Perplexity model — 'sonar' for research (default).

    Returns:
        Detailed, sourced research findings from the web.
    """
    if not PPLX_API_KEY:
        return "Error: PPLX_API_KEY not set in environment variables."

    try:
        resp = requests.post(
            "https://api.perplexity.ai/chat/completions",
            headers={
                "Authorization": f"Bearer {PPLX_API_KEY}",
                "Content-Type": "application/json",
            },
            json={
                "model": model,
                "messages": [{"role": "user", "content": query}],
                "return_citations": True,
            },
            timeout=90,
        )
        resp.raise_for_status()
        data     = resp.json()
        content  = data["choices"][0]["message"]["content"]
        citations = data.get("citations", [])
        if citations:
            content += "\n\nSources:\n" + "\n".join(f"- {url}" for url in citations[:10])
        return content
    except requests.exceptions.Timeout:
        return "Error: Perplexity API timed out. Try a more specific query."
    except Exception as exc:
        return f"Error calling Perplexity API: {exc}"


# ─────────────────────────────────────────────────────────────────────────────
# Word Report Generation
# ─────────────────────────────────────────────────────────────────────────────

@tool
def generate_docx_report(
    company_profile_json: str,
    employee_profile_json: str,
    industry_research_json: str,
    problems_research_json: str,
    solutions_research_json: str,
    output_path: str = "./reports/case_study.docx",
) -> str:
    """Generate a professional Word (.docx) marketing case study report.

    Combines all research into a formatted document ready for client delivery.

    Args:
        company_profile_json:    JSON string of company research data.
        employee_profile_json:   JSON string of employee/contact research data.
        industry_research_json:  JSON string of tech stack research.
        problems_research_json:  JSON string of identified pain points.
        solutions_research_json: JSON string of product solutions mapping.
        output_path:             File path for the .docx.

    Returns:
        Confirmation message with the path to the generated file.
    """
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        return "Error: python-docx not installed. Run: pip install python-docx"

    def _load(s) -> dict:
        if isinstance(s, dict):
            return s
        try:
            return json.loads(s)
        except Exception:
            return {}

    company   = _load(company_profile_json)
    employee  = _load(employee_profile_json)
    industry  = _load(industry_research_json)
    problems  = _load(problems_research_json)
    solutions = _load(solutions_research_json)

    doc = Document()

    # ── Cover Page ──────────────────────────────────────────────────────────
    t = doc.add_heading(f"Digital Transformation Strategy: {company.get('name', 'Company')}", 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Try to find the primary contact if multiple were found
    primary_contact_name = employee.get("full_name") or "Executive Team"
    primary_contact_title = employee.get("current_title") or ""
    
    if "leadership_team" in employee and employee["leadership_team"]:
        primary_contact_name = employee["leadership_team"][0].get("full_name", primary_contact_name)
        primary_contact_title = employee["leadership_team"][0].get("title", primary_contact_title)

    run = sub.add_run(
        f"Prepared for: {primary_contact_name} — {primary_contact_title}\n"
        f"Generated: {datetime.now().strftime('%B %d, %Y')}"
    )
    run.italic = True
    doc.add_page_break()

    # ── 1. Company Overview ─────────────────────────────────────────────────
    doc.add_heading("1. Company Overview", 1)
    _add_kv_table(doc, {
        "Company":   company.get("name", ""),
        "Website":   company.get("website", ""),
        "Industry":  company.get("industry", ""),
        "Sector":    company.get("sector", ""),
        "Employees": company.get("employee_count", ""),
        "Revenue":   company.get("estimated_revenue", ""),
        "HQ":        company.get("headquarters", ""),
    })
    if company.get("description"):
        doc.add_paragraph(company["description"])
    
    # Key Clients & Core Services
    if company.get("core_products_services"):
        doc.add_heading("Core Products & Services", 3)
        for item in company["core_products_services"]:
            doc.add_paragraph(str(item), style="List Bullet")
    if company.get("key_clients"):
        doc.add_heading("Key Clients", 3)
        for c in company["key_clients"]:
            doc.add_paragraph(str(c), style="List Bullet")


    # ── 2. Current Technology Landscape ─────────────────────────────────────
    doc.add_heading("2. Current Technology Landscape", 1)
    if industry.get("tech_stack_summary"):
        doc.add_paragraph(industry["tech_stack_summary"])
        
    doc.add_heading("Identified Tech Stack", 2)
    _add_kv_table(doc, {
        "ERP & Finance": ", ".join(industry.get("erp_finance", ["Not identified"])),
        "CRM & Sales": ", ".join(industry.get("crm", ["Not identified"])),
        "HRIS": ", ".join(industry.get("hris", ["Not identified"])),
        "Cloud Infra": ", ".join(industry.get("cloud_infrastructure", ["Not identified"])),
        "Collaboration": ", ".join(industry.get("collaboration_tools", ["Not identified"])),
        "Analytics": ", ".join(industry.get("marketing_analytics", ["Not identified"])),
    })
    
    if industry.get("tech_initiatives") and industry["tech_initiatives"] != ["Not identified"]:
        doc.add_heading("Known Tech Initiatives", 2)
        for init in industry["tech_initiatives"]:
            doc.add_paragraph(str(init), style="List Bullet")


    # ── 3. Operational Pain Points ──────────────────────────────────────────
    doc.add_heading("3. Operational Pain Points", 1)
    if problems.get("problems_summary"):
        doc.add_paragraph(problems["problems_summary"])
        
    if problems.get("problems"):
        for prob in problems["problems"]:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(prob.get("title", "")).bold = True
            if prob.get("description"):
                p.add_run(f": {prob['description']}")
            if prob.get("business_impact"):
                doc.add_paragraph(f"  → Impact: {prob['business_impact']}", style="List Bullet 2")


    # ── 4. Strategic NetSuite & AI Solutions ────────────────────────────────
    doc.add_heading("4. Strategic Solutions Mapping", 1)
    if solutions.get("solutions_summary"):
        doc.add_paragraph(solutions["solutions_summary"])
        
    if solutions.get("mapped_solutions"):
        for sol in solutions["mapped_solutions"]:
            doc.add_heading(sol.get("problem_title", "Solution"), 2)
            
            p = doc.add_paragraph()
            p.add_run("Recommended Technology: ").bold = True
            p.add_run(sol.get("recommended_solution", ""))
            
            doc.add_paragraph(sol.get("how_it_solves_it", ""))
            
            val = doc.add_paragraph(style="List Bullet")
            val.add_run("Business Value: ").bold = True
            val.add_run(sol.get("business_value", ""))


    # ── 5. Key Leadership ───────────────────────────────────────────────────
    doc.add_heading("5. Key Executive Contacts", 1)
    if "leadership_team" in employee and employee["leadership_team"]:
        for leader in employee["leadership_team"]:
            doc.add_paragraph(f"{leader.get('full_name', '')} — {leader.get('title', '')}", style="List Bullet")

    # ── Save ────────────────────────────────────────────────────────────────
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return f"✅ Report saved to: {output_path}"


def _add_kv_table(doc, data: dict) -> None:
    """Add a formatted two-column key/value table."""
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for key, value in data.items():
        if value:
            row = table.add_row()
            row.cells[0].text = key
            row.cells[1].text = str(value)
            for para in row.cells[0].paragraphs:
                for run in para.runs:
                    run.bold = True
    doc.add_paragraph()
